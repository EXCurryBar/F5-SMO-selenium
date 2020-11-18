import csv
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
import time
from datetime import datetime
import shutil

def words(ti, da, num):
    global word_nn, doc, t0, filecount
    # print(ti, ": ", da)
    if word_nn == 1:
        t0.cell(num, word_nn).text = da
    elif word_nn == 3:
        t0.cell(num, word_nn).text = da
    elif word_nn == 5:
        t0.cell(num, word_nn).text = da
    else:
        t0.cell(num, 6).text = da
    while True:
        try:
            doc.save("SMO_" + str(filecount) + ".docx")
            break
        except:
            pass


def paste(lst):
    sys_host, sys_sn, sys_uptime, sys_mem, sys_cpu, sys_ac, sys_nc, sys_tp, sys_log, sys_ntp, sys_snmp, sys_ucs, sys_qkview, sys_time, sys_cert, sys_ha, sys_ver = lst
    words("hostname", sys_host, 0)
    words("S/N", sys_sn, 1)
    words("uptime", sys_uptime, 2)
    words("Memory", sys_mem, 3)
    words("CPU", sys_cpu, 4)
    words("Active Connections", sys_ac, 5)
    words("New Connections", sys_nc, 6)
    words("Throughput", sys_tp, 7)
    words("Syslog", sys_log, 8)
    words("NTP", sys_ntp, 9)
    words("SNMP", sys_snmp, 10)
    words("Config Backup",sys_ucs,11)
    words("Qkview Backup",sys_qkview,12)
    words("Time", sys_time, 13)
    words("Certificate status",sys_cert,14)
    words("HA status", sys_ha, 15)
    words("Version", sys_ver, 16)


if __name__ == "__main__":
    global doc, t0, word_nn, filecount
    filecount = 0
    data_lst = []
    with open("data.csv", "r", encoding="utf-8") as csvf:
        data = csv.reader(csvf)
        for line in data:
            data_lst.append(line)

    for row in range(len(data_lst)):
        if row % 4 == 0:
            doc = Document('example.docx')
            doc.styles['Normal'].font.name = "Times New Roman"
            doc.styles['Normal'].font.size = Pt(10)
            t0 = doc.tables[0]
            word_nn = 1
            filecount += 1
        paste(data_lst[row])
        word_nn += 2
